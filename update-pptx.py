from pptx import Presentation
from pptx.util import Inches

prs = Presentation('./OLD_SYSTEM_PHP/archive/PPTx/Blood_Donation_Camp_System_Presentation.pptx')

# Replace text
for slide in prs.slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if "BLOOD DONATION" in run.text.upper() or "CAMP SYSTEM" in run.text.upper():
                    # Preserve case
                    if "BLOOD" in run.text:
                        run.text = run.text.replace("BLOOD DONATION CAMP SYSTEM", "DOTLIFE")
                    if "Blood" in run.text:
                        run.text = run.text.replace("Blood Donation Camp System", "DotLife")

# On slide 10, let's add the screenshots
slide_10 = prs.slides[9] # 0-indexed
# add landing.png and login.png
slide_10.shapes.add_picture('landing.png', Inches(0.5), Inches(2), height=Inches(2.5))
slide_10.shapes.add_picture('login.png', Inches(5.5), Inches(2), height=Inches(2.5))

prs.save('DotLife_Presentation.pptx')
