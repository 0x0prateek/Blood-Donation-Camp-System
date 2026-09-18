from pptx import Presentation

prs = Presentation('./OLD_SYSTEM_PHP/archive/PPTx/Blood_Donation_Camp_System_Presentation.pptx')

slide = prs.slides[0]
for shape in slide.shapes:
    if hasattr(shape, "text"):
        print("TEXT:", shape.text)
