from pptx import Presentation

prs = Presentation('./OLD_SYSTEM_PHP/archive/PPTx/Blood_Donation_Camp_System_Presentation.pptx')

print("Title:", prs.core_properties.title)

for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print("TEXT:", shape.text)
